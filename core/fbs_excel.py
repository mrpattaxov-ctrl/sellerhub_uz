# -*- coding: utf-8 -*-
"""Excel export/import for FBS/DBS SKU stocks (O6 «Обновить через файл»).

The file format mirrors Uzum's own stock template so a file downloaded
from EITHER platform imports into the other:

    Sheet «Список товаров», 9 columns —
      A SKU_ID · B SKU_Name · C Название товара · D Штрихкод ·
      E Идентификатор · F Подходит для схемы · G Остаток на складе ·
      H Продавать по FBS («Да») · I Продавать по DBS («Да»)

Only G/H/I are editable (green). Import is **header-driven** (columns are
located by keyword, with a fixed-position fallback) so column re-ordering
or a hand-made file still works.

Everything here is pure (bytes in / Workbook or dicts out) so the mutation
path is unit-tested before it ever reaches Uzum.
"""
from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STOCK_SHEET = "Список товаров"
INSTRUCT_SHEET = "Инструкция"

# Exact Uzum headers (bilingual) — kept verbatim for cross-platform parity.
HEADERS = [
    "SKU_ID",
    "SKU_Name",
    "Название товара\nTovarning nomi",
    "Штрихкод\nShtrixkod",
    "Идентификатор от продавца\nSotuvchidan identifikator",
    "Подходит для схемы работы\nIshlash sxemasiga mos keladi",
    "Остаток на складе, шт\nOmbordagi qoldiq, dona",
    "Продавать по FBS\nFBS bo‘yicha sotish\n\nНапишите — Да\nДа — deb yozing",
    "Продавать по DBS\nDBS bo‘yicha sotish\n\nНапишите — Да\nДа — deb yozing",
]

# 0-based positions used for the fixed-position fallback during import.
_POS_SKU_ID, _POS_AMOUNT, _POS_FBS, _POS_DBS = 0, 6, 7, 8

_YES = "Да"
_TRUE_TOKENS = {"да", "yes", "ha", "ha'", "true", "1", "+", "v", "✓"}

_GREEN = PatternFill("solid", fgColor="C6EFCE")   # editable columns
_GREY = PatternFill("solid", fgColor="F2F2F2")     # read-only columns


# ── helpers ──────────────────────────────────────────────────────────────
def _norm(s) -> str:
    """Lower-cased, newline-flattened header text for keyword matching."""
    return str(s or "").replace("\n", " ").strip().lower()


def _to_int(val):
    """Parse an int from int/float/str ('781715.0', '12', 12). None on fail."""
    if val is None:
        return None
    if isinstance(val, bool):
        return None
    if isinstance(val, int):
        return val
    if isinstance(val, float):
        return int(val)
    s = str(val).strip().replace(",", ".")
    if s == "":
        return None
    try:
        return int(float(s))
    except (TypeError, ValueError):
        return None


def _is_yes(val) -> bool:
    return _norm(val) in _TRUE_TOKENS


def _eligibility(fbs_allowed, dbs_allowed) -> str:
    if fbs_allowed and dbs_allowed:
        return "FBS/DBS"
    if fbs_allowed:
        return "FBS"
    if dbs_allowed:
        return "DBS"
    return ""


# ── export ───────────────────────────────────────────────────────────────
def build_stock_workbook(skus: list[dict]) -> openpyxl.Workbook:
    """Build the export workbook from GET /v2/fbs/sku/stocks rows.

    ``skus`` items use the Uzum field names: skuId, skuTitle, productTitle,
    barcode, amount, fbsAllowed, dbsAllowed, fbsLinked, dbsLinked,
    sellerSkuCode.
    """
    wb = openpyxl.Workbook()

    # Sheet 1 — short bilingual instructions.
    ins = wb.active
    ins.title = INSTRUCT_SHEET
    ins.column_dimensions["A"].width = 90
    lines = [
        "Как обновить остатки / Qoldiqni qanday yangilash",
        "",
        "1. Заполняйте только зелёные колонки: «Остаток», «Продавать по FBS», «Продавать по DBS».",
        "   Faqat yashil ustunlarni to‘ldiring: «Qoldiq», «FBS bo‘yicha sotish», «DBS bo‘yicha sotish».",
        "2. Чтобы включить схему — напишите «Да»; чтобы выключить — оставьте пусто.",
        "   Sxemani yoqish uchun «Да» deb yozing; o‘chirish uchun bo‘sh qoldiring.",
        "3. Не меняйте SKU_ID и штрихкод — по ним находится товар.",
        "   SKU_ID va shtrixkodni o‘zgartirmang — tovar shular bo‘yicha topiladi.",
        "4. Пустой остаток = строка пропускается (не обнуляется). 0 = обнулить.",
        "   Bo‘sh qoldiq = qator o‘tkazib yuboriladi (nollanmaydi). 0 = nolga tushirish.",
    ]
    for i, text in enumerate(lines, start=1):
        c = ins.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=13)

    # Sheet 2 — the data.
    ws = wb.create_sheet(STOCK_SHEET)
    header_font = Font(bold=True, color="1F2A44")
    header_align = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for j, head in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=j, value=head)
        c.font = header_font
        c.alignment = header_align
        # green = editable (G,H,I), grey = read-only
        c.fill = _GREEN if j in (_POS_AMOUNT + 1, _POS_FBS + 1, _POS_DBS + 1) else _GREY
    widths = [12, 26, 34, 18, 22, 16, 18, 18, 18]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"

    r = 2
    for s in (skus or []):
        if not isinstance(s, dict):
            continue
        sku_id = _to_int(s.get("skuId", s.get("id")))
        if sku_id is None:
            continue
        ws.cell(row=r, column=1, value=sku_id)
        ws.cell(row=r, column=2, value=s.get("sellerSkuCode") or s.get("skuTitle") or "")
        ws.cell(row=r, column=3, value=s.get("productTitle") or "")
        ws.cell(row=r, column=4, value=str(s.get("barcode") or ""))
        ws.cell(row=r, column=5, value="")
        ws.cell(row=r, column=6, value=_eligibility(s.get("fbsAllowed"), s.get("dbsAllowed")))
        ws.cell(row=r, column=7, value=_to_int(s.get("amount")) or 0)
        ws.cell(row=r, column=8, value=_YES if s.get("fbsLinked") else "")
        ws.cell(row=r, column=9, value=_YES if s.get("dbsLinked") else "")
        r += 1
    return ws.parent


def workbook_to_bytes(wb: openpyxl.Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── import ───────────────────────────────────────────────────────────────
def _pick_sheet(wb):
    if STOCK_SHEET in wb.sheetnames:
        return wb[STOCK_SHEET]
    # else: first sheet whose header row mentions SKU_ID
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if any("sku_id" in _norm(h).replace(" ", "") for h in first):
            return ws
    return wb.worksheets[0]


def _locate_columns(header_row) -> dict:
    """Map logical field → 0-based column index, by header keyword with a
    fixed-position fallback. Returns {sku_id, amount, fbs, dbs}."""
    idx = {"sku_id": None, "amount": None, "fbs": None, "dbs": None}
    for j, h in enumerate(header_row):
        n = _norm(h)
        flat = n.replace(" ", "")
        if idx["sku_id"] is None and ("sku_id" in flat or "skuid" in flat):
            idx["sku_id"] = j
        elif idx["amount"] is None and ("остаток" in n or "qoldiq" in n or "ombordagi" in n):
            idx["amount"] = j
        elif idx["fbs"] is None and "fbs" in n:
            idx["fbs"] = j
        elif idx["dbs"] is None and "dbs" in n:
            idx["dbs"] = j
    # fixed-position fallback (Uzum's layout) for anything still missing
    fallback = {"sku_id": _POS_SKU_ID, "amount": _POS_AMOUNT, "fbs": _POS_FBS, "dbs": _POS_DBS}
    for k, pos in fallback.items():
        if idx[k] is None and pos < len(header_row):
            idx[k] = pos
    return idx


def parse_stock_rows(file_bytes: bytes) -> tuple[list[dict], dict]:
    """Parse an uploaded stock xlsx → (rows, stats).

    Each row: {skuId, skuTitle, productTitle, barcode, amount, fbsLinked,
    dbsLinked}. A row is INCLUDED only when it has a usable int skuId AND a
    valid non-negative int amount — a blank/garbage amount is SKIPPED so a
    half-filled file never silently zeroes stock.

    stats: {total_data_rows, parsed, skipped_no_id, skipped_no_amount}.
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = _pick_sheet(wb)
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return [], {"total_data_rows": 0, "parsed": 0, "skipped_no_id": 0, "skipped_no_amount": 0}
    col = _locate_columns(all_rows[0])
    ci_id, ci_amt, ci_fbs, ci_dbs = col["sku_id"], col["amount"], col["fbs"], col["dbs"]

    def cell(row, ci):
        return row[ci] if (ci is not None and ci < len(row)) else None

    rows: list[dict] = []
    stats = {"total_data_rows": 0, "parsed": 0, "skipped_no_id": 0, "skipped_no_amount": 0}
    for raw in all_rows[1:]:
        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
            continue  # blank line
        stats["total_data_rows"] += 1
        sku_id = _to_int(cell(raw, ci_id))
        if sku_id is None:
            stats["skipped_no_id"] += 1
            continue
        amount = _to_int(cell(raw, ci_amt))
        if amount is None or amount < 0:
            stats["skipped_no_amount"] += 1
            continue
        rows.append({
            "skuId": sku_id,
            "skuTitle": (str(cell(raw, 1)) if cell(raw, 1) is not None else ""),
            "productTitle": (str(cell(raw, 2)) if cell(raw, 2) is not None else ""),
            "barcode": (str(cell(raw, 3)) if cell(raw, 3) is not None else ""),
            "amount": amount,
            "fbsLinked": _is_yes(cell(raw, ci_fbs)),
            "dbsLinked": _is_yes(cell(raw, ci_dbs)),
        })
        stats["parsed"] += 1
    return rows, stats


def diff_against_current(desired: list[dict], current: list[dict]) -> dict:
    """Compare parsed rows to the live GET stock; return only what changed.

    Returns {"changes": [...], "unchanged": int, "unknown": int}. A change
    carries old_* values for preview. ``unknown`` = file skuIds absent from
    the current eligible-SKU list (cannot be updated → skipped).
    """
    cur = {}
    for s in (current or []):
        if not isinstance(s, dict):
            continue
        sid = _to_int(s.get("skuId", s.get("id")))
        if sid is None:
            continue
        cur[sid] = {
            "amount": _to_int(s.get("amount")) or 0,
            "fbsLinked": bool(s.get("fbsLinked")),
            "dbsLinked": bool(s.get("dbsLinked")),
            "skuTitle": s.get("skuTitle") or s.get("sellerSkuCode") or "",
            "productTitle": s.get("productTitle") or "",
            "barcode": str(s.get("barcode") or ""),
        }

    changes, unchanged, unknown = [], 0, 0
    for d in desired:
        sid = d["skuId"]
        c = cur.get(sid)
        if c is None:
            unknown += 1
            continue
        if (d["amount"] == c["amount"] and d["fbsLinked"] == c["fbsLinked"]
                and d["dbsLinked"] == c["dbsLinked"]):
            unchanged += 1
            continue
        changes.append({
            "skuId": sid,
            # prefer the live descriptive fields (authoritative for the POST)
            "skuTitle": c["skuTitle"] or d["skuTitle"],
            "productTitle": c["productTitle"] or d["productTitle"],
            "barcode": c["barcode"] or d["barcode"],
            "amount": d["amount"],
            "fbsLinked": d["fbsLinked"],
            "dbsLinked": d["dbsLinked"],
            "old_amount": c["amount"],
            "old_fbsLinked": c["fbsLinked"],
            "old_dbsLinked": c["dbsLinked"],
        })
    return {"changes": changes, "unchanged": unchanged, "unknown": unknown}
