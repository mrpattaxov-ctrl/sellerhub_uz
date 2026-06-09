# -*- coding: utf-8 -*-
"""Tests for O6 Excel export/import (core.fbs_excel) — 2026-06-02.

The import feeds the FBS/DBS stock MUTATION, so per the action-endpoint
policy we prove the parse/diff BEFORE prod: columns are located correctly,
amounts coerce, blank amounts are SKIPPED (never silently zeroed), the
``Да``/empty flags round-trip, and the diff emits ONLY genuinely changed
rows. No Postgres, no Uzum HTTP.
"""
from __future__ import annotations

import io

import openpyxl

from core.fbs_excel import (
    HEADERS,
    STOCK_SHEET,
    build_stock_workbook,
    workbook_to_bytes,
    parse_stock_rows,
    diff_against_current,
)


def _wb_bytes(header, data_rows, sheet_title=STOCK_SHEET):
    """Build a minimal xlsx (one sheet) from explicit header + rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(list(header))
    for r in data_rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SAMPLE = [
    {"skuId": 781715, "sellerSkuCode": "LUX-A", "skuTitle": "A", "productTitle": "Prod A",
     "barcode": "1000007817150", "amount": 1, "fbsAllowed": True, "dbsAllowed": True,
     "fbsLinked": True, "dbsLinked": False},
    {"skuId": 781716, "sellerSkuCode": "LUX-B", "skuTitle": "B", "productTitle": "Prod B",
     "barcode": "1000007817167", "amount": 0, "fbsAllowed": True, "dbsAllowed": False,
     "fbsLinked": False, "dbsLinked": False},
]


class TestBuild:
    def test_two_sheets_and_headers(self):
        wb = build_stock_workbook(SAMPLE)
        assert STOCK_SHEET in wb.sheetnames
        ws = wb[STOCK_SHEET]
        header = [c.value for c in ws[1]]
        assert header == HEADERS  # exact Uzum parity

    def test_values_mapped(self):
        ws = build_stock_workbook(SAMPLE)[STOCK_SHEET]
        # row 2 = first SKU
        assert ws.cell(row=2, column=1).value == 781715          # SKU_ID int
        assert ws.cell(row=2, column=7).value == 1               # amount
        assert ws.cell(row=2, column=8).value == "Да"            # FBS linked
        assert ws.cell(row=2, column=9).value == ""              # DBS not linked
        assert ws.cell(row=2, column=6).value == "FBS/DBS"       # eligibility
        # row 3 = second SKU (FBS only, not linked)
        assert ws.cell(row=3, column=8).value == ""
        assert ws.cell(row=3, column=6).value == "FBS"

    def test_bytes_are_loadable(self):
        data = workbook_to_bytes(build_stock_workbook(SAMPLE))
        wb2 = openpyxl.load_workbook(io.BytesIO(data))
        assert STOCK_SHEET in wb2.sheetnames


class TestParse:
    def test_round_trip(self):
        data = workbook_to_bytes(build_stock_workbook(SAMPLE))
        rows, stats = parse_stock_rows(data)
        assert stats["parsed"] == 2
        by_id = {r["skuId"]: r for r in rows}
        assert by_id[781715]["amount"] == 1
        assert by_id[781715]["fbsLinked"] is True
        assert by_id[781715]["dbsLinked"] is False
        assert by_id[781716]["amount"] == 0           # explicit 0 kept
        assert by_id[781716]["fbsLinked"] is False

    def test_sku_id_float_string_coerced(self):
        data = _wb_bytes(HEADERS,
                         [["781715.0", "X", "P", "1000007817150", "", "FBS/DBS", "3.0", "Да", ""]])
        rows, _ = parse_stock_rows(data)
        assert rows[0]["skuId"] == 781715
        assert rows[0]["amount"] == 3
        assert rows[0]["fbsLinked"] is True

    def test_blank_amount_skipped_not_zeroed(self):
        # blank amount → row is SKIPPED (so a half-filled file can't zero stock)
        data = _wb_bytes(HEADERS, [
            [111, "A", "P", "100", "", "FBS/DBS", "", "Да", ""],     # blank amount
            [222, "B", "P", "200", "", "FBS/DBS", 5, "Да", ""],      # ok
        ])
        rows, stats = parse_stock_rows(data)
        assert [r["skuId"] for r in rows] == [222]
        assert stats["skipped_no_amount"] == 1

    def test_invalid_sku_id_skipped(self):
        data = _wb_bytes(HEADERS, [
            ["abc", "A", "P", "100", "", "FBS/DBS", 5, "Да", ""],    # bad id
            [333, "B", "P", "200", "", "FBS/DBS", 7, "", ""],        # ok
        ])
        rows, stats = parse_stock_rows(data)
        assert [r["skuId"] for r in rows] == [333]
        assert stats["skipped_no_id"] == 1

    def test_yes_token_variants(self):
        data = _wb_bytes(HEADERS, [
            [1, "A", "P", "1", "", "FBS/DBS", 1, "да", "ДА"],
            [2, "B", "P", "2", "", "FBS/DBS", 1, "yes", "нет"],
        ])
        rows, _ = parse_stock_rows(data)
        assert rows[0]["fbsLinked"] is True and rows[0]["dbsLinked"] is True
        assert rows[1]["fbsLinked"] is True and rows[1]["dbsLinked"] is False

    def test_columns_located_when_reordered(self):
        # DBS/FBS/amount/sku_id in a different order — header keywords win.
        header = ["Продавать по DBS", "Остаток на складе, шт", "Продавать по FBS",
                  "SKU_Name", "SKU_ID"]
        data = _wb_bytes(header, [["Да", 9, "", "name", 4040]])
        rows, _ = parse_stock_rows(data)
        assert rows[0]["skuId"] == 4040
        assert rows[0]["amount"] == 9
        assert rows[0]["fbsLinked"] is False
        assert rows[0]["dbsLinked"] is True

    def test_position_fallback_when_headers_unknown(self):
        # generic headers (no keywords) → fixed-position A/G/H/I fallback
        header = ["c0", "c1", "c2", "c3", "c4", "c5", "c6", "c7", "c8"]
        data = _wb_bytes(header,
                         [[555, "", "", "", "", "", 8, "Да", "Да"]])
        rows, _ = parse_stock_rows(data)
        assert rows[0]["skuId"] == 555
        assert rows[0]["amount"] == 8
        assert rows[0]["fbsLinked"] is True
        assert rows[0]["dbsLinked"] is True


class TestDiff:
    CURRENT = [
        {"skuId": 781715, "amount": 1, "fbsLinked": True, "dbsLinked": False,
         "skuTitle": "A", "productTitle": "Prod A", "barcode": "100"},
        {"skuId": 781716, "amount": 0, "fbsLinked": False, "dbsLinked": False,
         "skuTitle": "B", "productTitle": "Prod B", "barcode": "200"},
    ]

    def test_only_changed_rows_returned(self):
        desired = [
            {"skuId": 781715, "amount": 5, "fbsLinked": True, "dbsLinked": False,
             "skuTitle": "A", "productTitle": "Prod A", "barcode": "100"},   # amount 1→5
            {"skuId": 781716, "amount": 0, "fbsLinked": False, "dbsLinked": False,
             "skuTitle": "B", "productTitle": "Prod B", "barcode": "200"},   # unchanged
        ]
        res = diff_against_current(desired, self.CURRENT)
        assert [c["skuId"] for c in res["changes"]] == [781715]
        assert res["changes"][0]["amount"] == 5
        assert res["changes"][0]["old_amount"] == 1
        assert res["unchanged"] == 1
        assert res["unknown"] == 0

    def test_flag_change_detected(self):
        desired = [{"skuId": 781716, "amount": 0, "fbsLinked": True, "dbsLinked": False,
                    "skuTitle": "B", "productTitle": "P", "barcode": "200"}]  # FBS off→on
        res = diff_against_current(desired, self.CURRENT)
        assert len(res["changes"]) == 1
        assert res["changes"][0]["fbsLinked"] is True
        assert res["changes"][0]["old_fbsLinked"] is False

    def test_unknown_sku_skipped(self):
        desired = [{"skuId": 999999, "amount": 5, "fbsLinked": True, "dbsLinked": False,
                    "skuTitle": "Z", "productTitle": "P", "barcode": "9"}]
        res = diff_against_current(desired, self.CURRENT)
        assert res["changes"] == []
        assert res["unknown"] == 1

    def test_diff_uses_live_descriptive_fields(self):
        # POST body must carry the AUTHORITATIVE live skuTitle/barcode, not
        # whatever stale text the uploaded file had.
        desired = [{"skuId": 781715, "amount": 9, "fbsLinked": True, "dbsLinked": False,
                    "skuTitle": "STALE", "productTitle": "STALE", "barcode": "STALE"}]
        res = diff_against_current(desired, self.CURRENT)
        ch = res["changes"][0]
        assert ch["skuTitle"] == "A"
        assert ch["barcode"] == "100"
